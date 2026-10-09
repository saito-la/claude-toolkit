"""2つの docx を段落単位と文字単位で比べ、差分を Word の変更履歴（w:ins / w:del）として
新しい docx に書き出す。体裁は新版（revised）の docx を基にする。

usage: docx-track-changes.py <old.docx> <new.docx> <out.docx> [author]

両方を同じひな型（md2docx.py）で作っておくと、差分が本文の変更だけになる。
Word の比較コマンド（AppleScript の compare）が -1708 で動かない環境向けの代替。
段落は difflib で対応付け、対になった段落は文字単位の差分を w:ins / w:del にする。
3文字未満の変更されない部分は前後の変更にまとめ、履歴を読みやすくする。
"""
import copy
import difflib
import sys
from datetime import datetime, timezone

import docx
from docx.oxml.ns import qn
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

old_path, new_path, out_path = sys.argv[1:4]
author = sys.argv[4] if len(sys.argv) > 4 else 'revision'
date = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
rev_id = [1000]


def rev_attrs(el):
    rev_id[0] += 1
    el.set(qn('w:id'), str(rev_id[0]))
    el.set(qn('w:author'), author)
    el.set(qn('w:date'), date)
    return el


def make_run(text, rpr, deleted=False):
    r = etree.Element(qn('w:r'))
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = etree.SubElement(r, qn('w:delText') if deleted else qn('w:t'))
    t.text = text
    t.set(XML_SPACE, 'preserve')
    return r


def wrap(tag, run):
    el = rev_attrs(etree.Element(qn(tag)))
    el.append(run)
    return el


def first_rpr(p_el):
    for r in p_el.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        return copy.deepcopy(rpr) if rpr is not None else None
    return None


def clear_runs(p_el):
    for child in list(p_el):
        if child.tag != qn('w:pPr'):
            p_el.remove(child)


def mark_paragraph(p_el, tag):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        ppr = etree.Element(qn('w:pPr'))
        p_el.insert(0, ppr)
    rpr = ppr.find(qn('w:rPr'))
    if rpr is None:
        rpr = etree.SubElement(ppr, qn('w:rPr'))
    rpr.append(rev_attrs(etree.Element(qn(tag))))


def char_ops(a, b):
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    ops = [(t, a[i1:i2], b[j1:j2]) for t, i1, i2, j1, j2 in sm.get_opcodes()]
    # 短い equal を挟んだ変更は1つにまとめて読みやすくする
    merged = []
    for op in ops:
        if (op[0] == 'equal' and len(op[1]) < 3 and merged and merged[-1][0] != 'equal'):
            merged.append(op)
            continue
        if op[0] != 'equal' and len(merged) >= 2 and merged[-1][0] == 'equal' and len(merged[-1][1]) < 3 and merged[-2][0] != 'equal':
            eq = merged.pop()
            prev = merged.pop()
            merged.append(('replace', prev[1] + eq[1] + op[1], prev[2] + eq[2] + op[2]))
            continue
        if op[0] != 'equal' and merged and merged[-1][0] != 'equal':
            prev = merged.pop()
            merged.append(('replace', prev[1] + op[1], prev[2] + op[2]))
            continue
        merged.append(op)
    return merged


def rewrite_paragraph(p_el, old_text, new_text):
    rpr = first_rpr(p_el)
    clear_runs(p_el)
    for tag, a, b in char_ops(old_text, new_text):
        if tag == 'equal':
            p_el.append(make_run(a, rpr))
        else:
            if a:
                p_el.append(wrap('w:del', make_run(a, rpr, deleted=True)))
            if b:
                p_el.append(wrap('w:ins', make_run(b, rpr)))


def mark_inserted_paragraph(p_el):
    rpr = first_rpr(p_el)
    text = ''.join(t.text or '' for t in p_el.iter(qn('w:t')))
    clear_runs(p_el)
    p_el.append(wrap('w:ins', make_run(text, rpr)))
    mark_paragraph(p_el, 'w:ins')


def deleted_paragraph_from(old_p_el):
    p = copy.deepcopy(old_p_el)
    rpr = first_rpr(p)
    text = ''.join(t.text or '' for t in p.iter(qn('w:t')))
    clear_runs(p)
    p.append(wrap('w:del', make_run(text, rpr, deleted=True)))
    mark_paragraph(p, 'w:del')
    return p


old_doc = docx.Document(old_path)
new_doc = docx.Document(new_path)
old_ps = old_doc.paragraphs
new_ps = new_doc.paragraphs
a = [p.text for p in old_ps]
b = [p.text for p in new_ps]
body = new_doc.element.body
sect = body.find(qn('w:sectPr'))


def insert_deleted(old_idx_list, before_new_idx):
    anchor = new_ps[before_new_idx]._p if before_new_idx < len(new_ps) else sect
    for i in old_idx_list:
        dp = deleted_paragraph_from(old_ps[i]._p)
        if anchor is not None:
            anchor.addprevious(dp)
        else:
            body.append(dp)


stats = {'changed': 0, 'inserted': 0, 'deleted': 0}
for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
    if tag == 'equal':
        continue
    if tag == 'insert':
        for j in range(j1, j2):
            mark_inserted_paragraph(new_ps[j]._p)
            stats['inserted'] += 1
        continue
    if tag == 'delete':
        insert_deleted(range(i1, i2), j1)
        stats['deleted'] += i2 - i1
        continue
    # replace: 類似度で対にする
    olds = list(range(i1, i2))
    news = list(range(j1, j2))
    pairs = []
    while olds and news:
        best = None
        for i in olds:
            for j in news:
                r = difflib.SequenceMatcher(None, a[i], b[j], autojunk=False).ratio()
                if best is None or r > best[0]:
                    best = (r, i, j)
        if best[0] < 0.5:
            break
        pairs.append((best[1], best[2]))
        olds.remove(best[1])
        news.remove(best[2])
    for i, j in sorted(pairs):
        rewrite_paragraph(new_ps[j]._p, a[i], b[j])
        stats['changed'] += 1
    for j in news:
        mark_inserted_paragraph(new_ps[j]._p)
        stats['inserted'] += 1
    if olds:
        insert_deleted(olds, j1)
        stats['deleted'] += len(olds)

new_doc.save(out_path)
print(stats)
