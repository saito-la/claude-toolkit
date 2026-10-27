"""相手から届いた Word の雛型（抄録・申請書・様式）に、書式を変えずに文字を書き込む。

usage:
  docx-fill.py show <in.docx>                       段落と表のセルを番号付きで出す
  docx-fill.py set <in.docx> <out.docx> <edits.json>  edits の指定どおりに書き換えた docx を書き出す

edits.json は次の形の配列。番号は show が出す元の文書の番号で、書き換えの順番に左右されない。
  {"p": 6, "text": "本文…"}            段落 6 の文字を置き換える。"\\n" で段落を分け、同じ書式の段落を後ろに足す
  {"p": 8, "delete": true}            段落 8 を消す（雛型の説明文の段落など）
  {"cell": "t0r1c2", "text": "…"}     表 0 の 1 行 2 列のセルの文字を置き換える（セルの最初の段落の書式を使う）

置き換えた文字には、その段落で最も長い run の書式（書体・大きさ・太字）を使う。雛型の説明文が
「本文　MSゴシック10.5ポイント…」のように見出しの語と本文で大きさが違う run に分かれていても、本文の
書式が残る。段落の書式（揃え・字下げ・行間）は触らない。画像・図形・脚注は扱わない。
"""
import copy
import json
import re
import sys

import docx
from docx.oxml.ns import qn


def longest_run(p):
    runs = [r for r in p.runs if r.text]
    return max(runs, key=lambda r: len(r.text)) if runs else None


def set_text(p, text):
    src = longest_run(p)
    rpr = copy.deepcopy(src._r.find(qn('w:rPr'))) if src is not None and src._r.find(qn('w:rPr')) is not None else None
    for r in list(p._p.findall(qn('w:r'))) + list(p._p.findall(qn('w:hyperlink'))):
        p._p.remove(r)
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def fill_paragraph(p, text):
    lines = text.split('\n')
    set_text(p, lines[0])
    anchor = p._p
    for line in lines[1:]:
        new = copy.deepcopy(p._p)
        anchor.addnext(new)
        anchor = new
        set_text(docx.text.paragraph.Paragraph(new, p._parent), line)


def show(path):
    d = docx.Document(path)
    for i, p in enumerate(d.paragraphs):
        r = longest_run(p)
        size = r.font.size.pt if r is not None and r.font.size else None
        print(f'p{i}\t{p.style.name}\t{size or "-"}pt\t{p.text}')
    for t, table in enumerate(d.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                print(f't{t}r{ri}c{ci}\t{cell.text}')


def apply(src, out, edits_path):
    d = docx.Document(src)
    edits = json.load(open(edits_path, encoding='utf-8'))
    paras = list(d.paragraphs)
    tables = list(d.tables)
    # 段落を足し引きしても番号がずれないよう、元の段落の参照を先に取ってから書き換える
    for e in edits:
        if 'p' in e:
            p = paras[e['p']]
            if e.get('delete'):
                p._p.getparent().remove(p._p)
            else:
                fill_paragraph(p, e['text'])
        elif 'cell' in e:
            m = re.fullmatch(r't(\d+)r(\d+)c(\d+)', e['cell'])
            if not m:
                sys.exit(f'cell の書き方が違う：{e["cell"]}')
            cell = tables[int(m[1])].rows[int(m[2])].cells[int(m[3])]
            first, *rest = cell.paragraphs
            for q in rest:
                q._p.getparent().remove(q._p)
            fill_paragraph(first, e['text'])
        else:
            sys.exit(f'p か cell が要る：{e}')
    d.save(out)
    print(f'書き出した：{out}')


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == 'show':
        show(sys.argv[2])
    elif len(sys.argv) == 5 and sys.argv[1] == 'set':
        apply(*sys.argv[2:5])
    else:
        sys.exit(__doc__)
